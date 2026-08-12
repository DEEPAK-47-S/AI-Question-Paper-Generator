"""Private, on-demand PDF/DOCX/TXT examination paper exporters."""
from __future__ import annotations

import io
from collections import defaultdict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..models import GeneratedQuestionPaper


def _details(paper: GeneratedQuestionPaper) -> list[tuple[str, str]]:
    subject = paper.subject
    return [
        ("Department", subject.department.name), ("Subject", subject.name),
        ("Subject Code", subject.code), ("Semester", subject.semester.name),
        ("Regulation", subject.regulation.name), ("Examination", paper.exam_type),
        ("Academic Year", paper.academic_year), ("Date", paper.exam_date.isoformat() if paper.exam_date else "—"),
        ("Duration", f"{paper.duration_minutes // 60} Hours {paper.duration_minutes % 60 or ''}".strip()),
        ("Maximum Marks", str(paper.total_marks)),
    ]


def _sections(paper: GeneratedQuestionPaper):
    grouped = defaultdict(list)
    for question in paper.questions:
        grouped[question.section].append(question)
    section_config = {item["name"]: item for item in paper.settings_json.get("sections", [])}
    return [(name, questions, section_config.get(name, {})) for name, questions in grouped.items()]


class NumberedCanvasMixin:
    pass


def export_pdf(paper: GeneratedQuestionPaper, include_answers: bool = False) -> bytes:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("College", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=14, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#102B4E"), spaceAfter=3)
    subtitle = ParagraphStyle("Sub", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, alignment=TA_CENTER, spaceAfter=8)
    heading = ParagraphStyle("Part", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, alignment=TA_CENTER, textColor=colors.HexColor("#102B4E"), spaceBefore=10, spaceAfter=5)
    body = ParagraphStyle("Question", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=15, alignment=TA_LEFT, spaceAfter=6)
    marks = ParagraphStyle("Marks", parent=body, alignment=TA_RIGHT)
    answer_style = ParagraphStyle("Answer", parent=body, fontSize=9, leftIndent=8 * mm, textColor=colors.HexColor("#374151"), backColor=colors.HexColor("#F4F6F8"), borderPadding=5)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D1D5DB"))
        canvas.line(18 * mm, 14 * mm, 192 * mm, 14 * mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(18 * mm, 9 * mm, f"{paper.paper_code} • Confidential academic document")
        canvas.drawRightString(192 * mm, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=20 * mm, title=paper.title, author="NIET AI-QPG")
    story = [
        Paragraph("NEHRU INSTITUTE OF ENGINEERING AND TECHNOLOGY", title),
        Paragraph("Coimbatore, Tamil Nadu", subtitle),
        Paragraph(paper.subject.department.name.upper(), heading),
    ]
    details = _details(paper)
    rows = []
    for index in range(0, len(details), 2):
        left, right = details[index], details[index + 1]
        rows.append([Paragraph(f"<b>{left[0]}:</b> {left[1]}", body), Paragraph(f"<b>{right[0]}:</b> {right[1]}", body)])
    table = Table(rows, colWidths=[87 * mm, 87 * mm])
    table.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#CBD5E1")), ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#E2E8F0")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
    story.extend([table, Spacer(1, 6 * mm)])
    for name, questions, config in _sections(paper):
        count = len(questions)
        marks_each = questions[0].marks if questions else 0
        story.extend([
            Paragraph(name.upper(), heading),
            Paragraph(config.get("instructions", "Answer ALL questions"), subtitle),
            Paragraph(f"{count} × {marks_each} = {count * marks_each} Marks", subtitle),
        ])
        for question in questions:
            qtext = question.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            qtable = Table([[Paragraph(f"<b>{question.question_number}.</b> {qtext}", body), Paragraph(f"[{question.marks}]", marks)]], colWidths=[160 * mm, 14 * mm])
            qtable.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
            story.append(qtable)
            if question.alternative_text:
                alt = question.alternative_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.extend([Paragraph("<b>OR</b>", subtitle), Paragraph(alt, body)])
            if include_answers and question.answer_key:
                answer = question.answer_key.answer_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
                story.append(Paragraph(f"<b>Answer key:</b><br/>{answer}", answer_style))
    story.extend([Spacer(1, 8 * mm), Paragraph("— END OF QUESTION PAPER —", subtitle)])
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def export_docx(paper: GeneratedQuestionPaper, include_answers: bool = False) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.55)
    section.bottom_margin = Inches(.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NEHRU INSTITUTE OF ENGINEERING AND TECHNOLOGY")
    run.bold = True
    run.font.size = Pt(14)
    p = doc.add_paragraph("Coimbatore, Tamil Nadu")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(paper.subject.department.name.upper())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    details = _details(paper)
    for index in range(0, len(details), 2):
        cells = table.add_row().cells
        left, right = details[index], details[index + 1]
        cells[0].text = f"{left[0]}: {left[1]}"
        cells[1].text = f"{right[0]}: {right[1]}"
    for name, questions, config in _sections(paper):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(name.upper())
        run.bold = True
        p = doc.add_paragraph(config.get("instructions", "Answer ALL questions"))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(f"{len(questions)} × {questions[0].marks if questions else 0} = {sum(q.marks for q in questions)} Marks")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for question in questions:
            p = doc.add_paragraph(style="List Number")
            p.add_run(question.text)
            p.add_run(f"   [{question.marks}]").bold = True
            if question.alternative_text:
                p = doc.add_paragraph("OR")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph(question.alternative_text)
            if include_answers and question.answer_key:
                p = doc.add_paragraph()
                p.add_run("Answer key: ").bold = True
                p.add_run(question.answer_key.answer_text)
    p = doc.add_paragraph("END OF QUESTION PAPER")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.text = f"{paper.paper_code} • Confidential academic document"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def export_txt(paper: GeneratedQuestionPaper, include_answers: bool = False) -> bytes:
    lines = [
        "NEHRU INSTITUTE OF ENGINEERING AND TECHNOLOGY", "Coimbatore, Tamil Nadu",
        paper.subject.department.name.upper(), "", *[f"{key}: {value}" for key, value in _details(paper)], "",
    ]
    for name, questions, config in _sections(paper):
        lines += [name.upper(), config.get("instructions", "Answer ALL questions"), f"{len(questions)} x {questions[0].marks if questions else 0} = {sum(q.marks for q in questions)} Marks", ""]
        for question in questions:
            lines.append(f"{question.question_number}. {question.text} [{question.marks}]")
            if question.alternative_text:
                lines += ["   OR", f"   {question.alternative_text}"]
            if include_answers and question.answer_key:
                lines.append(f"   ANSWER KEY: {question.answer_key.answer_text}")
            lines.append("")
    lines += ["END OF QUESTION PAPER", paper.paper_code]
    return "\n".join(lines).encode("utf-8")
