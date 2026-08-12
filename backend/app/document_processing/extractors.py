"""Real document extraction for supported academic formats with page/sheet/slide provenance."""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ExtractedSection:
    text: str
    location: str


@dataclass
class ExtractionResult:
    sections: list[ExtractedSection] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(section.text for section in self.sections if section.text.strip())

    @property
    def page_count(self) -> int:
        return len(self.sections)


class ExtractionError(ValueError):
    pass


def _ocr_image(image) -> str:
    try:
        import pytesseract
        from ..config import get_settings
        return pytesseract.image_to_string(image, lang=get_settings().ocr_language).strip()
    except Exception as exc:
        raise ExtractionError("OCR could not read this image. Ensure Tesseract and the configured language pack are installed.") from exc


def _pdf(path: Path) -> ExtractionResult:
    import fitz
    from PIL import Image

    result = ExtractionResult(metadata={"format": "PDF", "ocr_pages": []})
    try:
        pdf = fitz.open(path)
    except Exception as exc:
        raise ExtractionError("The PDF is corrupted, encrypted, or unreadable.") from exc
    if pdf.needs_pass:
        pdf.close()
        raise ExtractionError("Password-protected PDFs are not supported. Upload an unlocked academic copy.")
    for index, page in enumerate(pdf):
        text = page.get_text("text").strip()
        if len(re.sub(r"\s+", "", text)) < 30:
            try:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr = _ocr_image(image)
                if len(ocr) > len(text):
                    text = ocr
                    result.metadata["ocr_pages"].append(index + 1)
            except ExtractionError:
                # Keep any native text and let the final empty-document check provide a useful error.
                pass
        result.sections.append(ExtractedSection(text=text, location=f"Page {index + 1}"))
    pdf.close()
    return result


def _docx(path: Path) -> ExtractionResult:
    from docx import Document
    try:
        doc = Document(path)
    except Exception as exc:
        raise ExtractionError("The Word document is corrupted or unreadable.") from exc
    blocks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            style = p.style.name if p.style else ""
            prefix = "# " if style.lower().startswith("heading") else ""
            blocks.append(prefix + p.text.strip())
    for number, table in enumerate(doc.tables, 1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks.append(f"Table {number}\n" + "\n".join(rows))
    return ExtractionResult([ExtractedSection("\n".join(blocks), "Document")], {"format": "DOCX", "tables": len(doc.tables)})


def _pptx(path: Path) -> ExtractionResult:
    from pptx import Presentation
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ExtractionError("The PowerPoint file is corrupted or unreadable.") from exc
    result = ExtractionResult(metadata={"format": "PPTX"})
    for index, slide in enumerate(presentation.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                lines.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes and notes.lower() != "slide notes":
                lines.append("Notes: " + notes)
        result.sections.append(ExtractedSection("\n".join(lines), f"Slide {index}"))
    return result


def _xlsx(path: Path) -> ExtractionResult:
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("The Excel workbook is corrupted or unreadable.") from exc
    result = ExtractionResult(metadata={"format": "XLSX", "sheets": workbook.sheetnames})
    for sheet in workbook.worksheets:
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                lines.append(" | ".join(values))
        result.sections.append(ExtractedSection("\n".join(lines), f"Sheet {sheet.title}"))
    workbook.close()
    return result


def _csv(path: Path) -> ExtractionResult:
    raw = path.read_bytes()
    text = _decode_text(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [" | ".join(row) for row in csv.reader(io.StringIO(text), dialect)]
    return ExtractionResult([ExtractedSection("\n".join(rows), "CSV")], {"format": "CSV"})


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Text encoding could not be detected.")


def _text(path: Path) -> ExtractionResult:
    text = _decode_text(path.read_bytes())
    return ExtractionResult([ExtractedSection(text, "Text")], {"format": "TXT"})


def _image(path: Path) -> ExtractionResult:
    from PIL import Image
    try:
        with Image.open(path) as image:
            text = _ocr_image(image.convert("RGB"))
            metadata = {"format": image.format or "IMAGE", "dimensions": list(image.size), "ocr_pages": [1]}
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("The image is corrupted or unsupported.") from exc
    return ExtractionResult([ExtractedSection(text, "Image 1")], metadata)


def _html(path: Path) -> ExtractionResult:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(_decode_text(path.read_bytes()), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()
    text = "\n".join(part.strip() for part in soup.get_text("\n").splitlines() if part.strip())
    return ExtractionResult([ExtractedSection(text, "Web page")], {"format": "HTML", "title": soup.title.string.strip() if soup.title and soup.title.string else ""})


def _legacy_office(path: Path, extension: str) -> ExtractionResult:
    """Convert DOC/PPT/XLS using LibreOffice when present, then use the native parser."""
    if not shutil.which("libreoffice"):
        raise ExtractionError(f"Legacy {extension.upper()} extraction requires LibreOffice. Save as the modern Office format or install LibreOffice.")
    target_ext = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}[extension]
    with tempfile.TemporaryDirectory() as temp:
        process = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", target_ext[1:], "--outdir", temp, str(path)],
            capture_output=True, text=True, timeout=60, check=False,
        )
        converted = Path(temp) / (path.stem + target_ext)
        if process.returncode or not converted.exists():
            raise ExtractionError(f"LibreOffice could not convert the {extension.upper()} document.")
        return extract_document(converted, target_ext)


SUPPORTED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx", ".csv", ".txt",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".html", ".htm",
}


def extract_document(path: Path, extension: str | None = None) -> ExtractionResult:
    extension = (extension or path.suffix).lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type: {extension or 'unknown'}")
    if path.stat().st_size == 0:
        raise ExtractionError("The uploaded document is empty.")
    parsers = {
        ".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx,
        ".csv": _csv, ".txt": _text, ".html": _html, ".htm": _html,
        ".png": _image, ".jpg": _image, ".jpeg": _image, ".tif": _image,
        ".tiff": _image, ".bmp": _image, ".webp": _image,
    }
    result = _legacy_office(path, extension) if extension in {".doc", ".ppt", ".xls"} else parsers[extension](path)
    if len(re.sub(r"\s+", "", result.text)) < 10:
        raise ExtractionError("No usable academic text was found. The scan may be unreadable or OCR may be unavailable.")
    return result
