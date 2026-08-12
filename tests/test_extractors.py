from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from backend.app.document_processing import extractors
from backend.app.document_processing.extractors import extract_document
from backend.app.document_processing.pipeline import extract_structure


def test_pdf_extraction(tmp_path):
    path = tmp_path / "sample.pdf"
    pdf = fitz.open(); page = pdf.new_page(); page.insert_text((72, 72), "UNIT I - MATRICES\n1. Eigenvalues and eigenvectors are prescribed topics."); pdf.save(path); pdf.close()
    result = extract_document(path)
    assert "Eigenvalues" in result.text
    assert result.sections[0].location == "Page 1"


def test_docx_table_extraction(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document(); doc.add_heading("UNIT II - CONTROL FLOW", 1); doc.add_paragraph("1. Conditional statements and loops")
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "Topic"; table.cell(0, 1).text = "Hours"; doc.save(path)
    result = extract_document(path)
    assert "Conditional statements" in result.text and "Topic | Hours" in result.text


def test_excel_extraction(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Unit III"; ws.append(["Topic", "Pointers"]); wb.save(path)
    result = extract_document(path)
    assert "Topic | Pointers" in result.text
    assert result.sections[0].location == "Sheet Unit III"


def test_powerpoint_extraction(tmp_path):
    path = tmp_path / "sample.pptx"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = "UNIT IV - FILES"; slide.placeholders[1].text = "Text and binary file handling"; deck.save(path)
    result = extract_document(path)
    assert "binary file handling" in result.text
    assert result.sections[0].location == "Slide 1"


def test_image_ocr_pipeline(tmp_path, monkeypatch):
    from PIL import Image
    path = tmp_path / "scan.png"
    Image.new("RGB", (300, 100), "white").save(path)
    monkeypatch.setattr(extractors, "_ocr_image", lambda image: "UNIT V - HASHING\nCollision resolution and hash functions")
    result = extract_document(path)
    assert "Collision resolution" in result.text
    assert result.metadata["ocr_pages"] == [1]


def test_syllabus_unit_structure():
    text = "UNIT I - BASICS\n1. Variables\n2. Data Types\nUNIT II - LOOPS\n1. For Loop\n2. While Loop"
    units = extract_structure(text)
    assert [unit["unit"] for unit in units] == ["UNIT I", "UNIT II"]
    assert units[0]["topics"][0]["topic"] == "Variables"
